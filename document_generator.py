"""
OmniScribe - Document Generator
Produces high-fidelity ethnographic dossiers in DOCX (Word), PDF, Markdown, TXT, SRT, and VTT.
Features:
- Pure Native Language Script (original script, e.g. Devanagari)
- Pure, genuine verbatim English Translation
- Structured speaker turns (Interviewer vs Artisan / Interviewee)
- Complete Artisan Dossiers aggregating cross-folder sections with side-by-side bilingual comparison
- Master Combined Dossiers for entire folders/batches
"""

import html
import json
import os
import re
import urllib.parse
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import docx
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    HRFlowable,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from artisan_organizer import extract_artisan_and_section, group_results_by_artisan
from transcription_engine import TranscriptionResult, TranscriptionSegment

# Register Windows Indic & Unicode fonts for PDF rendering
DEFAULT_PDF_FONT = "Helvetica"
try:
    if os.path.exists("C:/Windows/Fonts/Nirmala.ttc"):
        pdfmetrics.registerFont(TTFont("Nirmala", "C:/Windows/Fonts/Nirmala.ttc", subfontIndex=0))
        DEFAULT_PDF_FONT = "Nirmala"
    elif os.path.exists("C:/Windows/Fonts/segoeui.ttf"):
        pdfmetrics.registerFont(TTFont("SegoeUI", "C:/Windows/Fonts/segoeui.ttf"))
        DEFAULT_PDF_FONT = "SegoeUI"
    elif os.path.exists("C:/Windows/Fonts/arial.ttf"):
        pdfmetrics.registerFont(TTFont("Arial", "C:/Windows/Fonts/arial.ttf"))
        DEFAULT_PDF_FONT = "Arial"
except Exception:
    DEFAULT_PDF_FONT = "Helvetica"


def _set_cell_background(cell, hex_color: str):
    """Sets background shading of a docx table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)


def _set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    """Sets cell internal padding."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def parse_dialogue_turns(text: str) -> List[Tuple[str, str]]:
    """
    Parses raw transcription text into dialogue speaker turns.
    Returns list of (speaker_name, dialogue_text).
    """
    if not text or not text.strip():
        return []

    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
    turns: List[Tuple[str, str]] = []

    # Regex: Look for speaker headers like "Interviewer:", "Deepak Prajapati:", "प्रश्नकर्ता:", etc.
    speaker_regex = re.compile(r'^([A-Za-z0-9\u0900-\u097F\u0980-\u09FF\s\-\.\'\(\)]+?)\s*:\s*(.*)$')

    current_speaker = ""
    current_content = []

    for line in lines:
        m = speaker_regex.match(line)
        if m and not line.lower().startswith(("http://", "https://", "time:", "timestamp:", "chunk:")):
            if current_speaker or current_content:
                turns.append((current_speaker, " ".join(current_content).strip()))
                current_content = []
            current_speaker = m.group(1).strip()
            rest = m.group(2).strip()
            if rest:
                current_content.append(rest)
        else:
            current_content.append(line)

    if current_speaker or current_content:
        turns.append((current_speaker, " ".join(current_content).strip()))

    # If no speaker labels detected at all, return whole text as single turn
    if not turns:
        turns.append(("", text.strip()))

    return turns


def is_interviewer_speaker(speaker: str) -> bool:
    """Checks if a speaker name represents the interviewer."""
    s = speaker.lower()
    return any(k in s for k in ["interviewer", "question", "प्रश्नकर्ता", "सवाल", "investigator", "researcher", "host"])


def render_dialogue_to_docx(container, text: str, default_artisan: str = "Artisan"):
    """
    Renders structured dialogue turns into a docx cell or document
    with clear visual color hierarchy for Interviewer vs Artisan.
    """
    turns = parse_dialogue_turns(text)
    if not turns:
        p = container.add_paragraph() if hasattr(container, "add_paragraph") else container.paragraphs[0]
        p.add_run("(No spoken dialogue detected)").italic = True
        return

    first = True
    for speaker, content in turns:
        if first and hasattr(container, "paragraphs") and len(container.paragraphs) > 0 and not container.paragraphs[0].text:
            p = container.paragraphs[0]
            first = False
        else:
            p = container.add_paragraph()
            first = False

        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15

        if speaker:
            run_spk = p.add_run(f"{speaker}: ")
            run_spk.font.name = "Segoe UI"
            run_spk.font.size = Pt(10)
            run_spk.font.bold = True

            if is_interviewer_speaker(speaker):
                run_spk.font.color.rgb = RGBColor(29, 78, 216)  # Bold Indigo Blue
            else:
                run_spk.font.color.rgb = RGBColor(5, 150, 105)  # Bold Emerald Green

        run_txt = p.add_run(content)
        run_txt.font.name = "Segoe UI"
        run_txt.font.size = Pt(10)
        run_txt.font.color.rgb = RGBColor(30, 41, 59)


def render_dialogue_to_pdf_paragraphs(
    text: str,
    default_artisan: str,
    body_style: ParagraphStyle
) -> List[Paragraph]:
    """Converts dialogue turns into styled ReportLab Paragraph flowables."""
    turns = parse_dialogue_turns(text)
    if not turns:
        return [Paragraph("<i>(No spoken dialogue detected)</i>", body_style)]

    paragraphs = []
    for speaker, content in turns:
        safe_content = html.escape(content).replace("\n", "<br/>")
        if speaker:
            safe_speaker = html.escape(speaker)
            color = "#1d4ed8" if is_interviewer_speaker(speaker) else "#059669"
            html_line = f'<b><font color="{color}">{safe_speaker}:</font></b> <font color="#1e293b">{safe_content}</font>'
        else:
            html_line = f'<font color="#1e293b">{safe_content}</font>'

        paragraphs.append(Paragraph(html_line, body_style))
        paragraphs.append(Spacer(1, 4))

    return paragraphs


class DocumentGenerator:
    """Creates structured individual, artisan-centric, and master batch documents."""

    @staticmethod
    def format_srt_time(seconds: float) -> str:
        """Converts seconds into SRT timestamp 00:00:00,000."""
        hrs = int(seconds // 3600)
        mins = int((seconds % 3600) // 60)
        secs = int(seconds % 60)
        millis = int(round((seconds - int(seconds)) * 1000))
        return f"{hrs:02d}:{mins:02d}:{secs:02d},{millis:03d}"

    @staticmethod
    def format_vtt_time(seconds: float) -> str:
        """Converts seconds into VTT timestamp 00:00:00.000."""
        hrs = int(seconds // 3600)
        mins = int((seconds % 3600) // 60)
        secs = int(seconds % 60)
        millis = int(round((seconds - int(seconds)) * 1000))
        return f"{hrs:02d}:{mins:02d}:{secs:02d}.{millis:03d}"

    @staticmethod
    def format_duration_static(seconds: float) -> str:
        hours = int(seconds // 3600)
        minutes = int((seconds % 3600) // 60)
        secs = int(seconds % 60)
        if hours > 0:
            return f"{hours}h {minutes}m {secs}s"
        elif minutes > 0:
            return f"{minutes}m {secs}s"
        else:
            return f"{secs}s"

    # -------------------------------------------------------------
    # SUBTITLES GENERATOR (.srt and .vtt)
    # -------------------------------------------------------------
    def generate_subtitles(
        self,
        result: TranscriptionResult,
        output_dir: Path
    ) -> dict:
        """Generates SRT and VTT subtitle files for Native and English."""
        files_created = {}

        for kind, attr, ext_label in [
            ("Native", "native", "native"),
            ("English", "english", "english")
        ]:
            # SRT
            srt_path = output_dir / f"Subtitles_{kind}.srt"
            with open(srt_path, "w", encoding="utf-8") as f:
                if result.segments:
                    for idx, seg in enumerate(result.segments, 1):
                        txt = getattr(seg, attr, "") or seg.native
                        if txt.strip():
                            f.write(f"{idx}\n")
                            f.write(f"{self.format_srt_time(seg.start)} --> {self.format_srt_time(seg.end)}\n")
                            f.write(f"{txt.strip()}\n\n")
                else:
                    # Write whole dialogue as single subtitle block
                    text_content = result.native_text if kind == "Native" else result.english_text
                    if text_content.strip():
                        f.write(f"1\n00:00:00,000 --> {self.format_srt_time(result.duration_seconds)}\n{text_content.strip()}\n\n")
            files_created[f"srt_{ext_label}"] = srt_path

            # VTT
            vtt_path = output_dir / f"Subtitles_{kind}.vtt"
            with open(vtt_path, "w", encoding="utf-8") as f:
                f.write("WEBVTT\n\n")
                if result.segments:
                    for idx, seg in enumerate(result.segments, 1):
                        txt = getattr(seg, attr, "") or seg.native
                        if txt.strip():
                            f.write(f"{self.format_vtt_time(seg.start)} --> {self.format_vtt_time(seg.end)}\n")
                            f.write(f"{txt.strip()}\n\n")
                else:
                    text_content = result.native_text if kind == "Native" else result.english_text
                    if text_content.strip():
                        f.write(f"{self.format_vtt_time(0)} --> {self.format_vtt_time(result.duration_seconds)}\n{text_content.strip()}\n\n")
            files_created[f"vtt_{ext_label}"] = vtt_path

        return files_created

    # -------------------------------------------------------------
    # INDIVIDUAL DOCX GENERATION
    # -------------------------------------------------------------
    def generate_docx(
        self,
        result: TranscriptionResult,
        output_path: Path
    ) -> Path:
        """
        Generates a clean, professional Microsoft Word (.docx) document
        containing pure Native Script and pure English Translation dialogue turns.
        """
        doc = Document()

        for s in doc.sections:
            s.top_margin = Inches(0.8)
            s.bottom_margin = Inches(0.8)
            s.left_margin = Inches(0.8)
            s.right_margin = Inches(0.8)

        # Title
        p_title = doc.add_paragraph()
        run_title = p_title.add_run("🎙️ OmniScribe Transcription Dossier")
        run_title.font.name = "Segoe UI"
        run_title.font.size = Pt(20)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(15, 23, 42)
        p_title.paragraph_format.space_after = Pt(2)

        # Subtitle
        p_sub = doc.add_paragraph()
        run_sub = p_sub.add_run(
            f"Artisan: {result.artisan_name or 'Artisan'} • Section: {result.section_name or 'General'} • Generated on {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        )
        run_sub.font.name = "Segoe UI"
        run_sub.font.size = Pt(10)
        run_sub.font.color.rgb = RGBColor(100, 116, 139)
        p_sub.paragraph_format.space_after = Pt(12)

        # Metadata Table
        table = doc.add_table(rows=5, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        meta_rows = [
            ("👤 Artisan / Interviewee", result.artisan_name or "Recognized from Recording"),
            ("📌 Section / Topic", result.section_name or "General Interview"),
            ("📁 Source Audio File", result.filename),
            ("⏱️ Total Duration", result.duration_formatted),
            ("🌐 Spoken Language", f"{result.language_detected} (AI Confidence: {int(result.confidence * 100)}%)")
        ]

        for i, (label, val) in enumerate(meta_rows):
            cell_lbl = table.rows[i].cells[0]
            cell_val = table.rows[i].cells[1]
            cell_lbl.width = Inches(2.2)
            cell_val.width = Inches(4.5)

            _set_cell_background(cell_lbl, "F8FAFC")
            _set_cell_background(cell_val, "FFFFFF")
            _set_cell_margins(cell_lbl, top=80, bottom=80, left=140, right=140)
            _set_cell_margins(cell_val, top=80, bottom=80, left=140, right=140)

            p_l = cell_lbl.paragraphs[0]
            r_l = p_l.add_run(label)
            r_l.font.name = "Segoe UI"
            r_l.font.bold = True
            r_l.font.size = Pt(9.5)
            r_l.font.color.rgb = RGBColor(51, 65, 85)

            p_v = cell_val.paragraphs[0]
            r_v = p_v.add_run(val)
            r_v.font.name = "Segoe UI"
            r_v.font.size = Pt(9.5)
            r_v.font.color.rgb = RGBColor(15, 23, 42)

        doc.add_paragraph().paragraph_format.space_after = Pt(14)

        # 1. Native Transcription Section
        h1 = doc.add_heading(level=1)
        r_h1 = h1.add_run("1. 🗣️ Native Language Transcription (Original Script)")
        r_h1.font.name = "Segoe UI"
        r_h1.font.size = Pt(13)
        r_h1.font.bold = True
        r_h1.font.color.rgb = RGBColor(5, 150, 105)

        callout1 = doc.add_table(rows=1, cols=1)
        callout1.alignment = WD_TABLE_ALIGNMENT.CENTER
        c1 = callout1.rows[0].cells[0]
        c1.width = Inches(6.7)
        _set_cell_background(c1, "F0FDF4")
        _set_cell_margins(c1, top=140, bottom=140, left=160, right=160)
        render_dialogue_to_docx(c1, result.native_text, result.artisan_name)

        doc.add_paragraph().paragraph_format.space_after = Pt(14)

        # 2. English Translation Section
        h2 = doc.add_heading(level=1)
        r_h2 = h2.add_run("2. 🌐 English Translation (Genuine Raw Dialogue)")
        r_h2.font.name = "Segoe UI"
        r_h2.font.size = Pt(13)
        r_h2.font.bold = True
        r_h2.font.color.rgb = RGBColor(29, 78, 216)

        callout2 = doc.add_table(rows=1, cols=1)
        callout2.alignment = WD_TABLE_ALIGNMENT.CENTER
        c2 = callout2.rows[0].cells[0]
        c2.width = Inches(6.7)
        _set_cell_background(c2, "EEF2FF")
        _set_cell_margins(c2, top=140, bottom=140, left=160, right=160)
        render_dialogue_to_docx(c2, result.english_text, result.artisan_name)

        doc.save(str(output_path))
        return output_path

    # -------------------------------------------------------------
    # INDIVIDUAL PDF GENERATION
    # -------------------------------------------------------------
    def generate_pdf(
        self,
        result: TranscriptionResult,
        output_path: Path
    ) -> Path:
        """Generates a clean PDF report using ReportLab with Indic script support."""
        doc = SimpleDocTemplate(
            str(output_path),
            pagesize=letter,
            rightMargin=40,
            leftMargin=40,
            topMargin=40,
            bottomMargin=40
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Normal"],
            fontName=DEFAULT_PDF_FONT,
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#0f172a"),
            fontBold=True,
            spaceAfter=4
        )

        sub_style = ParagraphStyle(
            "DocSub",
            parent=styles["Normal"],
            fontName=DEFAULT_PDF_FONT,
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor("#64748b"),
            spaceAfter=12
        )

        h1_style = ParagraphStyle(
            "Heading1Style",
            parent=styles["Heading1"],
            fontName=DEFAULT_PDF_FONT,
            fontSize=12,
            leading=16,
            textColor=colors.HexColor("#0f172a"),
            fontBold=True,
            spaceBefore=10,
            spaceAfter=6
        )

        body_style = ParagraphStyle(
            "BodyStyle",
            parent=styles["Normal"],
            fontName=DEFAULT_PDF_FONT,
            fontSize=9.5,
            leading=14,
            textColor=colors.HexColor("#1e293b")
        )

        story = []
        story.append(Paragraph("🎙️ OmniScribe Transcription Dossier", title_style))
        story.append(Paragraph(
            f"Artisan: {html.escape(result.artisan_name or 'Artisan')} • Section: {html.escape(result.section_name or 'General')} • {datetime.now().strftime('%Y-%m-%d %H:%M')}",
            sub_style
        ))

        # Metadata Table
        meta_data = [
            [Paragraph("<b>Artisan / Interviewee:</b>", body_style), Paragraph(html.escape(result.artisan_name or "Recognized"), body_style)],
            [Paragraph("<b>Section / Topic:</b>", body_style), Paragraph(html.escape(result.section_name or "General"), body_style)],
            [Paragraph("<b>Audio File:</b>", body_style), Paragraph(html.escape(result.filename), body_style)],
            [Paragraph("<b>Duration:</b>", body_style), Paragraph(result.duration_formatted, body_style)],
            [Paragraph("<b>Language:</b>", body_style), Paragraph(f"{html.escape(result.language_detected)} (Conf: {int(result.confidence*100)}%)", body_style)]
        ]
        meta_table = Table(meta_data, colWidths=[140, 390])
        meta_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor("#f8fafc")),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(meta_table)
        story.append(Spacer(1, 14))

        # 1. Native Script
        story.append(Paragraph("1. 🗣️ Native Language Transcription (Original Script)", h1_style))
        story.append(Spacer(1, 6))
        for p in render_dialogue_to_pdf_paragraphs(result.native_text, result.artisan_name, body_style):
            story.append(p)
        story.append(Spacer(1, 14))

        # 2. English Translation
        story.append(Paragraph("2. 🌐 English Translation (Genuine Raw Dialogue)", h1_style))
        story.append(Spacer(1, 6))
        for p in render_dialogue_to_pdf_paragraphs(result.english_text, result.artisan_name, body_style):
            story.append(p)

        try:
            doc.build(story)
        except Exception:
            # Fallback font
            try:
                title_style.fontName = "Helvetica"
                sub_style.fontName = "Helvetica"
                h1_style.fontName = "Helvetica"
                body_style.fontName = "Helvetica"
                doc.build(story)
            except Exception:
                pass

        return output_path

    # -------------------------------------------------------------
    # MARKDOWN & TXT GENERATION
    # -------------------------------------------------------------
    def generate_markdown(
        self,
        result: TranscriptionResult,
        output_path: Path
    ) -> Path:
        """Generates a clean Markdown report."""
        md = f"""# 🎙️ OmniScribe Transcription Dossier

**Artisan / Interviewee:** `{result.artisan_name or 'Artisan'}`  
**Section / Topic:** `{result.section_name or 'General Interview'}`  
**Source Audio File:** `{result.filename}`  
**Duration:** `{result.duration_formatted}` ({result.duration_seconds:.1f}s)  
**Spoken Language:** `{result.language_detected}` (Confidence: {int(result.confidence * 100)}%)  
**Generated At:** `{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}`  

---

## 1. 🗣️ Native Language Transcription (Original Script)

{result.native_text or "*(No spoken dialogue detected)*"}

---

## 2. 🌐 English Translation (Genuine Raw Dialogue)

{result.english_text or "*(No spoken dialogue detected)*"}
"""
        output_path.write_text(md, encoding="utf-8")
        return output_path

    def generate_raw_txt(
        self,
        result: TranscriptionResult,
        output_dir: Path
    ) -> dict:
        """Generates clean standalone TXT files for Native and English."""
        files = {}
        # 1. Native
        p_nat = output_dir / "01_Native_Transcription.txt"
        p_nat.write_text(result.native_text, encoding="utf-8")
        files["txt_native"] = p_nat

        # 2. English
        p_eng = output_dir / "02_English_Translation.txt"
        p_eng.write_text(result.english_text, encoding="utf-8")
        files["txt_english"] = p_eng

        return files

    # -------------------------------------------------------------
    # CONSOLIDATED ARTISAN COMPLETE DOSSIER
    # (Aggregates cross-folder sections with side-by-side bilingual view)
    # -------------------------------------------------------------
    def generate_artisan_dossier(
        self,
        artisan_name: str,
        sections: List[TranscriptionResult],
        output_dir: Path
    ) -> dict:
        """
        Creates an aggregated complete dossier for an artisan across all interview sections.
        Features side-by-side bilingual comparison (Native on left, English on right).
        Generates:
        1. <Artisan>_Complete_Dossier.docx
        2. <Artisan>_Complete_Dossier.pdf
        3. <Artisan>_Complete_Dossier.md
        4. <Artisan>_Complete_Dossier.txt
        """
        output_dir.mkdir(parents=True, exist_ok=True)
        artisan_slug = re.sub(r'[^\w\-_\.]', '_', artisan_name).strip('_')
        created = {}

        total_sections = len(sections)
        total_duration = sum(s.duration_seconds for s in sections)
        total_duration_str = self.format_duration_static(total_duration)

        # 1. Word DOCX
        docx_path = output_dir / f"{artisan_slug}_Complete_Dossier.docx"
        doc = Document()
        for s in doc.sections:
            s.top_margin = Inches(0.8)
            s.bottom_margin = Inches(0.8)
            s.left_margin = Inches(0.8)
            s.right_margin = Inches(0.8)

        # Header Title
        p_title = doc.add_paragraph()
        run_title = p_title.add_run(f"👤 Artisan Complete Dossier: {artisan_name}")
        run_title.font.name = "Segoe UI"
        run_title.font.size = Pt(22)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(15, 23, 42)

        p_sub = doc.add_paragraph()
        p_sub.add_run(
            f"Consolidated Cross-Section Field Interviews • Total Sections: {total_sections} • Cumulative Duration: {total_duration_str} • Compiled on {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        )
        p_sub.paragraph_format.space_after = Pt(14)

        # Table of Sections
        doc.add_heading("📑 Table of Interview Sections", level=1)
        toc_table = doc.add_table(rows=total_sections + 1, cols=4)
        toc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for col_i, h in enumerate(["Section", "Source File", "Duration", "Detected Language"]):
            c = toc_table.rows[0].cells[col_i]
            _set_cell_background(c, "E2E8F0")
            _set_cell_margins(c, top=80, bottom=80, left=120, right=120)
            c.paragraphs[0].add_run(h).font.bold = True

        for idx, sec in enumerate(sections, 1):
            row = toc_table.rows[idx]
            for c in row.cells:
                _set_cell_margins(c, top=60, bottom=60, left=120, right=120)
            row.cells[0].paragraphs[0].add_run(sec.section_name or f"Section {idx}")
            row.cells[1].paragraphs[0].add_run(sec.filename)
            row.cells[2].paragraphs[0].add_run(sec.duration_formatted)
            row.cells[3].paragraphs[0].add_run(sec.language_detected)

        doc.add_page_break()

        # Each Section: Side-by-Side Table (Native on Left, English on Right)
        for idx, sec in enumerate(sections, 1):
            h_sec = doc.add_heading(f"{sec.section_name or f'Section {idx}'}: {sec.filename}", level=1)
            h_sec.paragraph_format.space_before = Pt(14)
            h_sec.paragraph_format.space_after = Pt(4)

            p_m = doc.add_paragraph()
            p_m.add_run(f"⏱️ Duration: {sec.duration_formatted} | 🌐 Language: {sec.language_detected} | 🤖 Engine: {sec.engine_used.upper()}")
            p_m.paragraph_format.space_after = Pt(10)

            # Side-by-Side 2-Column Table
            side_table = doc.add_table(rows=2, cols=2)
            side_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            side_table.autofit = False

            # Headers
            c_h_left = side_table.rows[0].cells[0]
            c_h_right = side_table.rows[0].cells[1]
            c_h_left.width = Inches(3.35)
            c_h_right.width = Inches(3.35)

            _set_cell_background(c_h_left, "F0FDF4")
            _set_cell_background(c_h_right, "EEF2FF")
            _set_cell_margins(c_h_left, top=100, bottom=100, left=140, right=140)
            _set_cell_margins(c_h_right, top=100, bottom=100, left=140, right=140)

            r_hl = c_h_left.paragraphs[0].add_run(f"🗣️ Native Script ({sec.language_detected})")
            r_hl.font.name = "Segoe UI"
            r_hl.font.bold = True
            r_hl.font.color.rgb = RGBColor(5, 150, 105)

            r_hr = c_h_right.paragraphs[0].add_run("🌐 English Translation")
            r_hr.font.name = "Segoe UI"
            r_hr.font.bold = True
            r_hr.font.color.rgb = RGBColor(29, 78, 216)

            # Dialogue Bodies
            c_b_left = side_table.rows[1].cells[0]
            c_b_right = side_table.rows[1].cells[1]
            c_b_left.width = Inches(3.35)
            c_b_right.width = Inches(3.35)

            _set_cell_background(c_b_left, "FFFFFF")
            _set_cell_background(c_b_right, "FFFFFF")
            _set_cell_margins(c_b_left, top=120, bottom=120, left=140, right=140)
            _set_cell_margins(c_b_right, top=120, bottom=120, left=140, right=140)

            render_dialogue_to_docx(c_b_left, sec.native_text, sec.artisan_name)
            render_dialogue_to_docx(c_b_right, sec.english_text, sec.artisan_name)

            if idx < total_sections:
                doc.add_page_break()

        doc.save(str(docx_path))
        created["dossier_docx"] = docx_path

        # 2. Markdown
        md_path = output_dir / f"{artisan_slug}_Complete_Dossier.md"
        md_text = f"# 👤 Artisan Complete Dossier: {artisan_name}\n\n"
        md_text += f"- **Total Sections:** {total_sections}\n"
        md_text += f"- **Cumulative Duration:** {total_duration_str}\n"
        md_text += f"- **Compiled At:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n"
        md_text += "## 📑 Table of Sections\n\n"
        md_text += "| Section | Source File | Duration | Language |\n|---|---|---|---|\n"
        for s in sections:
            md_text += f"| {s.section_name} | `{s.filename}` | {s.duration_formatted} | {s.language_detected} |\n"
        md_text += "\n---\n\n"

        for idx, sec in enumerate(sections, 1):
            md_text += f"## {sec.section_name}: `{sec.filename}`\n"
            md_text += f"**Duration:** {sec.duration_formatted} | **Language:** {sec.language_detected}\n\n"
            md_text += "### 🗣️ Native Language Dialogue (Original Script)\n\n"
            md_text += f"{sec.native_text or '*(No dialogue)*'}\n\n"
            md_text += "### 🌐 English Translation (Genuine Raw Dialogue)\n\n"
            md_text += f"{sec.english_text or '*(No dialogue)*'}\n\n"
            md_text += "---\n\n"

        md_path.write_text(md_text, encoding="utf-8")
        created["dossier_md"] = md_path

        # 3. Plain TXT
        txt_path = output_dir / f"{artisan_slug}_Complete_Dossier.txt"
        txt_c = f"======================================================================\n"
        txt_c += f"ARTISAN COMPLETE DOSSIER: {artisan_name.upper()}\n"
        txt_c += f"Total Sections: {total_sections} | Total Duration: {total_duration_str}\n"
        txt_c += f"Compiled: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
        txt_c += f"======================================================================\n\n"
        for idx, sec in enumerate(sections, 1):
            txt_c += f"======================================================================\n"
            txt_c += f"SECTION {idx:02d}: {sec.section_name.upper()} ({sec.filename})\n"
            txt_c += f"Duration: {sec.duration_formatted} | Language: {sec.language_detected}\n"
            txt_c += f"======================================================================\n\n"
            txt_c += f"[NATIVE SCRIPT]:\n{sec.native_text}\n\n"
            txt_c += f"[ENGLISH TRANSLATION]:\n{sec.english_text}\n\n\n"
        txt_path.write_text(txt_c, encoding="utf-8")
        created["dossier_txt"] = txt_path

        # 4. PDF (ReportLab)
        pdf_path = output_dir / f"{artisan_slug}_Complete_Dossier.pdf"
        try:
            pdf_doc = SimpleDocTemplate(
                str(pdf_path),
                pagesize=letter,
                leftMargin=36,
                rightMargin=36,
                topMargin=36,
                bottomMargin=36
            )
            pdf_story = []
            styles = getSampleStyleSheet()

            p_h1 = ParagraphStyle("PH1", parent=styles["Heading1"], fontName=DEFAULT_PDF_FONT, fontSize=16, leading=20, textColor=colors.HexColor("#0f172a"))
            p_h2 = ParagraphStyle("PH2", parent=styles["Heading2"], fontName=DEFAULT_PDF_FONT, fontSize=11, leading=15, textColor=colors.HexColor("#1e293b"), spaceBefore=10)
            p_sub_style = ParagraphStyle("PSub", parent=styles["Normal"], fontName=DEFAULT_PDF_FONT, fontSize=9, leading=13, textColor=colors.HexColor("#64748b"))
            p_cell = ParagraphStyle("PCell", parent=styles["Normal"], fontName=DEFAULT_PDF_FONT, fontSize=8.5, leading=12)

            pdf_story.append(Paragraph(f"<b>👤 Artisan Complete Dossier: {html.escape(artisan_name)}</b>", p_h1))
            pdf_story.append(Paragraph(f"Sections: {total_sections} | Total Duration: {total_duration_str} | Compiled: {datetime.now().strftime('%Y-%m-%d %H:%M')}", p_sub_style))
            pdf_story.append(Spacer(1, 12))

            for idx, sec in enumerate(sections, 1):
                pdf_story.append(Paragraph(f"<b>📌 Section: {html.escape(sec.section_name)} ({html.escape(sec.filename)})</b>", p_h2))
                pdf_story.append(Paragraph(f"Duration: {sec.duration_formatted} | Language: {html.escape(sec.language_detected)}", p_sub_style))
                pdf_story.append(Spacer(1, 6))

                pdf_story.append(Paragraph(f"<b><font color='#059669'>🗣️ 1. Native Script ({html.escape(sec.language_detected)}):</font></b>", p_h2))
                for p in render_dialogue_to_pdf_paragraphs(sec.native_text, artisan_name, p_cell):
                    pdf_story.append(p)
                pdf_story.append(Spacer(1, 8))

                pdf_story.append(Paragraph("<b><font color='#1d4ed8'>🌐 2. English Translation (Verbatim):</font></b>", p_h2))
                for p in render_dialogue_to_pdf_paragraphs(sec.english_text, artisan_name, p_cell):
                    pdf_story.append(p)
                pdf_story.append(Spacer(1, 14))

            pdf_doc.build(pdf_story)
            created["dossier_pdf"] = pdf_path
        except Exception as e:
            print(f"[DocumentGenerator] PDF build error for {artisan_name}: {e}")

        return created

    # -------------------------------------------------------------
    # MASTER COMBINED BATCH DOCUMENT GENERATION
    # -------------------------------------------------------------
    def generate_master_combined_documents(
        self,
        batch_name: str,
        results: List[TranscriptionResult],
        output_dir: Path
    ) -> dict:
        """
        Combines all audio transcriptions from an entire folder/batch into:
        1. Master_Combined_Transcript.docx
        2. Master_Combined_Transcript.pdf
        3. Master_Combined_Transcript.md
        4. Master_Combined_Transcript.txt
        5. Master_Batch_Summary.json
        6. Artisans_Dossiers/ folder containing per-artisan dossiers.
        """
        output_dir.mkdir(parents=True, exist_ok=True)
        created = {}

        total_files = len(results)
        total_duration = sum(r.duration_seconds for r in results)
        total_duration_str = self.format_duration_static(total_duration)

        # Group by Artisan
        artisans_map = group_results_by_artisan(results)

        # Generate individual dossiers for each Artisan
        artisans_dir = output_dir / "Artisans_Dossiers"
        artisans_dir.mkdir(parents=True, exist_ok=True)
        artisans_summary_data = {}

        for artisan_name, art_sections in artisans_map.items():
            dossier_files = self.generate_artisan_dossier(artisan_name, art_sections, artisans_dir)
            art_duration = sum(s.duration_seconds for s in art_sections)
            art_slug = re.sub(r'[^\w\-_\.]', '_', artisan_name).strip('_')

            artisans_summary_data[artisan_name] = {
                "artisan_name": artisan_name,
                "sections_count": len(art_sections),
                "total_duration_seconds": art_duration,
                "total_duration_formatted": self.format_duration_static(art_duration),
                "sections": [
                    {
                        "section_name": s.section_name,
                        "filename": s.filename,
                        "duration_formatted": s.duration_formatted,
                        "language": s.language_detected
                    }
                    for s in art_sections
                ],
                "downloads": {
                    "docx": f"/api/download?path={urllib.parse.quote(str(artisans_dir / f'{art_slug}_Complete_Dossier.docx'))}",
                    "pdf": f"/api/download?path={urllib.parse.quote(str(artisans_dir / f'{art_slug}_Complete_Dossier.pdf'))}",
                    "md": f"/api/download?path={urllib.parse.quote(str(artisans_dir / f'{art_slug}_Complete_Dossier.md'))}",
                    "txt": f"/api/download?path={urllib.parse.quote(str(artisans_dir / f'{art_slug}_Complete_Dossier.txt'))}"
                }
            }

        # 1. Master DOCX
        docx_path = output_dir / "Master_Combined_Transcript.docx"
        doc = Document()
        for s in doc.sections:
            s.top_margin = Inches(0.8)
            s.bottom_margin = Inches(0.8)
            s.left_margin = Inches(0.8)
            s.right_margin = Inches(0.8)

        # Title
        p_title = doc.add_paragraph()
        run_title = p_title.add_run(f"📚 Master Transcription Dossier: {batch_name}")
        run_title.font.name = "Segoe UI"
        run_title.font.size = Pt(22)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(15, 23, 42)

        p_sub = doc.add_paragraph()
        p_sub.add_run(
            f"Batch Summary • Total Files: {total_files} • Detected Artisans: {len(artisans_map)} • Cumulative Duration: {total_duration_str} • Compiled on {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        )
        p_sub.paragraph_format.space_after = Pt(14)

        # Table of Artisans & Sections
        doc.add_heading("👥 Artisans & Sections Overview", level=1)
        art_table = doc.add_table(rows=len(artisans_map) + 1, cols=4)
        art_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for col_i, h in enumerate(["Artisan Name", "Sections Count", "Cumulative Duration", "Sections Included"]):
            c = art_table.rows[0].cells[col_i]
            _set_cell_background(c, "E2E8F0")
            c.paragraphs[0].add_run(h).font.bold = True

        for idx, (artisan_name, art_sections) in enumerate(artisans_map.items(), 1):
            row = art_table.rows[idx]
            row.cells[0].paragraphs[0].add_run(artisan_name)
            row.cells[1].paragraphs[0].add_run(str(len(art_sections)))
            art_dur = sum(s.duration_seconds for s in art_sections)
            row.cells[2].paragraphs[0].add_run(self.format_duration_static(art_dur))
            sec_names = ", ".join(s.section_name for s in art_sections)
            row.cells[3].paragraphs[0].add_run(sec_names)

        doc.add_page_break()

        # Each Audio Entry organized by Artisan
        for artisan_name, art_sections in artisans_map.items():
            doc.add_heading(f"👤 Artisan: {artisan_name}", level=1)

            for sec in art_sections:
                doc.add_heading(f"📌 {sec.section_name}: {sec.filename}", level=2)
                p_m = doc.add_paragraph()
                p_m.add_run(f"Duration: {sec.duration_formatted} | Language: {sec.language_detected} | Engine: {sec.engine_used.upper()}")
                p_m.paragraph_format.space_after = Pt(6)

                # Native
                doc.add_heading("🗣️ Native Script (Original Language)", level=3)
                c1 = doc.add_table(rows=1, cols=1).rows[0].cells[0]
                c1.width = Inches(6.7)
                _set_cell_background(c1, "F0FDF4")
                _set_cell_margins(c1)
                render_dialogue_to_docx(c1, sec.native_text, artisan_name)

                # English
                doc.add_heading("🌐 English Translation (Genuine Raw Dialogue)", level=3)
                c2 = doc.add_table(rows=1, cols=1).rows[0].cells[0]
                c2.width = Inches(6.7)
                _set_cell_background(c2, "EEF2FF")
                _set_cell_margins(c2)
                render_dialogue_to_docx(c2, sec.english_text, artisan_name)

                doc.add_paragraph().paragraph_format.space_after = Pt(10)

            doc.add_page_break()

        doc.save(str(docx_path))
        created["master_docx"] = docx_path

        # 2. Master Markdown
        md_path = output_dir / "Master_Combined_Transcript.md"
        md_text = f"# 📚 Master Audio Transcription Dossier: {batch_name}\n\n"
        md_text += f"- **Total Files:** {total_files}\n"
        md_text += f"- **Total Artisans:** {len(artisans_map)}\n"
        md_text += f"- **Total Duration:** {total_duration_str}\n"
        md_text += f"- **Compiled At:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n---\n\n"

        for artisan_name, art_sections in artisans_map.items():
            md_text += f"# 👤 Artisan: {artisan_name}\n\n"
            for sec in art_sections:
                md_text += f"## {sec.section_name}: `{sec.filename}`\n"
                md_text += f"**Duration:** {sec.duration_formatted} | **Language:** {sec.language_detected}\n\n"
                md_text += "### 🗣️ Native Script\n\n"
                md_text += f"{sec.native_text or '*(No dialogue)*'}\n\n"
                md_text += "### 🌐 English Translation\n\n"
                md_text += f"{sec.english_text or '*(No dialogue)*'}\n\n"
                md_text += "---\n\n"

        md_path.write_text(md_text, encoding="utf-8")
        created["master_md"] = md_path

        # 3. Master TXT
        txt_path = output_dir / "Master_Combined_Transcript.txt"
        txt_content = f"===========================================================\n"
        txt_content += f"MASTER AUDIO TRANSCRIPTION DOSSIER: {batch_name}\n"
        txt_content += f"Total Files: {total_files} | Total Duration: {total_duration_str}\n"
        txt_content += f"===========================================================\n\n"
        for artisan_name, art_sections in artisans_map.items():
            txt_content += f"###########################################################\n"
            txt_content += f"ARTISAN: {artisan_name.upper()}\n"
            txt_content += f"###########################################################\n\n"
            for sec in art_sections:
                txt_content += f"--- [{sec.section_name}: {sec.filename}] ---\n"
                txt_content += f"Duration: {sec.duration_formatted} | Language: {sec.language_detected}\n\n"
                txt_content += f"[NATIVE SCRIPT]:\n{sec.native_text}\n\n"
                txt_content += f"[ENGLISH TRANSLATION]:\n{sec.english_text}\n\n"
                txt_content += f"-----------------------------------------------------------\n\n"
        txt_path.write_text(txt_content, encoding="utf-8")
        created["master_txt"] = txt_path

        # 4. Master JSON
        json_path = output_dir / "Master_Batch_Summary.json"
        batch_dict = {
            "batch_name": batch_name,
            "total_files": total_files,
            "total_duration_seconds": total_duration,
            "total_duration_formatted": total_duration_str,
            "total_artisans": len(artisans_map),
            "artisans": artisans_summary_data,
            "compiled_at": datetime.now().isoformat(),
            "results": [r.to_dict() for r in results]
        }
        json_path.write_text(json.dumps(batch_dict, indent=2, ensure_ascii=False), encoding="utf-8")
        created["master_json"] = json_path

        # 5. Master PDF
        pdf_path = output_dir / "Master_Combined_Transcript.pdf"
        try:
            pdf_doc = SimpleDocTemplate(str(pdf_path), pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
            pdf_story = []
            styles = getSampleStyleSheet()

            m_h1 = ParagraphStyle("MH1", parent=styles["Heading1"], fontName=DEFAULT_PDF_FONT, fontSize=16, leading=20)
            m_h2 = ParagraphStyle("MH2", parent=styles["Heading2"], fontName=DEFAULT_PDF_FONT, fontSize=12, leading=16, spaceBefore=10)
            m_h3 = ParagraphStyle("MH3", parent=styles["Heading3"], fontName=DEFAULT_PDF_FONT, fontSize=10, leading=14, spaceBefore=6)
            m_norm = ParagraphStyle("MNorm", parent=styles["Normal"], fontName=DEFAULT_PDF_FONT, fontSize=9, leading=13)

            pdf_story.append(Paragraph(f"<b>Master Transcription Dossier: {html.escape(batch_name)}</b>", m_h1))
            pdf_story.append(Paragraph(f"Files: {total_files} | Artisans: {len(artisans_map)} | Cumulative Duration: {total_duration_str}", m_norm))
            pdf_story.append(Spacer(1, 14))

            for artisan_name, art_sections in artisans_map.items():
                pdf_story.append(Paragraph(f"<b>👤 Artisan: {html.escape(artisan_name)}</b>", m_h2))
                for sec in art_sections:
                    pdf_story.append(Paragraph(f"<b>📌 {html.escape(sec.section_name)}: {html.escape(sec.filename)}</b> ({sec.duration_formatted})", m_h3))
                    pdf_story.append(Paragraph("<b>🗣️ Native Script:</b>", m_norm))
                    for p in render_dialogue_to_pdf_paragraphs(sec.native_text, artisan_name, m_norm):
                        pdf_story.append(p)
                    pdf_story.append(Spacer(1, 6))
                    pdf_story.append(Paragraph("<b>🌐 English Translation:</b>", m_norm))
                    for p in render_dialogue_to_pdf_paragraphs(sec.english_text, artisan_name, m_norm):
                        pdf_story.append(p)
                    pdf_story.append(Spacer(1, 12))

            pdf_doc.build(pdf_story)
            created["master_pdf"] = pdf_path
        except Exception:
            pass

        return created
